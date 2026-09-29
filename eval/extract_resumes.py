"""提取 eval 简历文本并做轻量脱敏.

输入: ~/uploads/<uuid>/Rxx.docx|pdf  (从 Paul Mac 上传的原始简历)
输出:
  eval/raw_resumes/Rxx.<ext>   原始文件归档(带正确文件名)
  eval/resumes_text/Rxx.txt     脱敏后纯文本,供 eval 脚本使用

脱敏规则(2026-09-28, Paul 确认"没必要"做姓名脱敏, 仅程序化去除联系方式):
  - 邮箱 -> [EMAIL]
  - 电话(美式 10 位及常见分隔符, 国际区号) -> [PHONE]
  - 姓名保留原文(学生自愿提供给 Paul 的简历; eval 需调用 Anthropic API, 已告知 Paul)
"""
import os
import re
import shutil
import sys

HOME = os.path.expanduser("~")
UPLOADS = os.path.join(HOME, "uploads")
EVAL = os.path.join(HOME, "workspace", "resume-compass", "eval")
RAW = os.path.join(EVAL, "raw_resumes")
TEXT = os.path.join(EVAL, "resumes_text")
os.makedirs(RAW, exist_ok=True)
os.makedirs(TEXT, exist_ok=True)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(
    r"(?:\+\d{1,3}[\s\-.]?)?"          # 可选国际区号
    r"(?:\(?\d{3}\)?[\s\-.]?)"          # 区号
    r"\d{3}[\s\-.]?\d{4}"              # 本地号码(美式)
    r"|(?:\(86\)|\+86)?[\s\-.]?1[3-9]\d[\s\-.]?\d{4}[\s\-.]?\d{4}"  # 中国大陆手机
    r"|(?:\(852\)|\+852)[\s\-.]?\d{4}[\s\-.]?\d{4}"                 # 香港(必须带852区号,避免误伤年份)
)


def redact(text: str) -> str:
    text = EMAIL_RE.sub("[EMAIL]", text)
    text = PHONE_RE.sub("[PHONE]", text)
    return text


def extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def main() -> None:
    # 找到所有上传文件,按 Rxx 归一化
    found: dict[str, str] = {}
    for root, _dirs, files in os.walk(UPLOADS):
        for f in files:
            m = re.match(r"(R\d{2})\.(docx|pdf)$", f)
            if m:
                found[m.group(1)] = os.path.join(root, f)
    print(f"找到 {len(found)} 个简历文件")
    missing = [f"R{i:02d}" for i in range(1, 31) if f"R{i:02d}" not in found]
    if missing:
        print(f"缺失: {missing}")
        sys.exit(1)

    for rid in sorted(found):
        src = found[rid]
        ext = os.path.splitext(src)[1]
        dst_raw = os.path.join(RAW, f"{rid}{ext}")
        shutil.copy2(src, dst_raw)
        raw_text = extract_docx(src) if ext == ".docx" else extract_pdf(src)
        clean = redact(raw_text)
        with open(os.path.join(TEXT, f"{rid}.txt"), "w", encoding="utf-8") as fh:
            fh.write(clean)
        n_email = clean.count("[EMAIL]")
        n_phone = clean.count("[PHONE]")
        print(f"{rid}: {len(clean)} 字符, 脱敏 email×{n_email} phone×{n_phone}")

    print("完成")


if __name__ == "__main__":
    main()
